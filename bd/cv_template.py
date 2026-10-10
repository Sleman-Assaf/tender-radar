# bd/cv_template.py
# Path: TenderRadar/bd/cv_template.py
#
# Generates a formatted CV DOCX in the standard AJi CV template layout from an
# Engineer's structured data. Requires python-docx (`pip install python-docx`).
#
# Template section order (from CV_Templete_1.docx):
#   Name (bold) / Current Position / Profession / Nationality
#   [Date of Birth / Years with Firm when present]
#   Membership in Professional Societies (bullets)
#   Countries of Work Experience
#   Synopsis
#   Education
#   Employment Record (employer, position, period + responsibility bullets)
#   Selected Experience (project, country, year / position / client / description)
#   PROFESSIONAL COURSES / CERTIFICATIONS
#   Languages table (Language | Reading | Speaking | Writing)

import io
import logging

logger = logging.getLogger('tenders')


def generate_cv_docx(engineer, photo_bytes: bytes = None) -> bytes:
    """Build the formatted CV and return the .docx bytes. When photo_bytes is
    None, the engineer's stored photo is used if present.
    Raises RuntimeError with a clear message if python-docx is missing."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise RuntimeError(
            'python-docx is not installed on the server. '
            'Run: pip install python-docx'
        )

    doc = Document()

    # ── Base style: Arial per AJi brand ─────────────────────────────────────
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10.5)

    def para(text='', bold=False, size=None, space_after=4, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if align:
            p.alignment = align
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        return p

    def label_value(label, value):
        if not value:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f'{label}: ')
        r.bold = True
        p.add_run(str(value))

    def section(title, upper=False, underline=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(title.upper() if upper else title)
        r.bold = True
        r.underline = underline
        r.font.size = Pt(12)
        return p

    def bullet(text, bold=False):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold = bold

    # ── Photo (top, like the AJi template) ──────────────────────────────────
    if photo_bytes is None and getattr(engineer, 'photo', None):
        try:
            with engineer.photo.open('rb') as ph:
                photo_bytes = ph.read()
        except Exception:
            photo_bytes = None
    if photo_bytes:
        def _add_photo(data):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            p.add_run().add_picture(io.BytesIO(data), height=Inches(1.4))
        try:
            _add_photo(photo_bytes)          # PNG/JPEG — sniffed by header
        except Exception:
            # Probably an EMF/WMF metafile — convert then retry
            try:
                from .photo_extract import _emf_to_png
                png = _emf_to_png(photo_bytes)
                if png:
                    _add_photo(png)
                else:
                    logger.warning('[bd.cv_template] photo skipped — EMF/WMF '
                                   'with no converter available (install Pillow).')
            except Exception as exc:
                logger.warning(f'[bd.cv_template] photo could not be embedded: {exc}')

    # ── Header block ─────────────────────────────────────────────────────────
    para(engineer.full_name, bold=True, size=16, space_after=2,
         align=WD_ALIGN_PARAGRAPH.CENTER)
    if engineer.current_position:
        para(engineer.current_position, size=12, space_after=10,
             align=WD_ALIGN_PARAGRAPH.CENTER)

    label_value('Profession',  engineer.profession)
    label_value('Nationality', engineer.nationality)
    label_value('Date of Birth',   engineer.date_of_birth)
    label_value('Years with Firm', engineer.years_with_firm)

    # ── Memberships ─────────────────────────────────────────────────────────
    memberships = engineer.memberships_list
    if memberships:
        section('Membership in Professional Societies')
        for m in memberships:
            bullet(m)

    label_value('Countries of Work Experience', engineer.countries_experience)

    # ── Synopsis ─────────────────────────────────────────────────────────────
    if engineer.synopsis:
        section('Synopsis')
        for chunk in engineer.synopsis.split('\n\n'):
            if chunk.strip():
                para(chunk.strip(), space_after=6)

    # ── Education ─────────────────────────────────────────────────────────────
    education = engineer.education
    if education:
        section('Education')
        for e in education:
            head = ', '.join(x for x in
                             [e.get('institution', ''), e.get('country', '')] if x)
            years = e.get('years', '')
            if years:
                head = f'{head}, {years}' if head else years
            if head:
                para(head, bold=True, space_after=1)
            if e.get('degree'):
                para(e['degree'], space_after=6)

    # ── Employment Record ────────────────────────────────────────────────────
    employment = engineer.employment
    if employment:
        section('Employment Record')
        for job in employment:
            head = ', '.join(x for x in
                             [job.get('employer', ''), job.get('position', '')] if x)
            period = job.get('period', '')
            if period:
                head = f'{head}, {period}' if head else period
            if head:
                para(head, bold=True, space_after=2)
            resp = job.get('responsibilities') or []
            if resp:
                para('Responsibilities:', bold=True, space_after=2)
                for rline in resp:
                    if rline:
                        bullet(rline)

    # ── Selected Experience ──────────────────────────────────────────────────
    experience = engineer.experience
    if experience:
        p = section('Selected Experience', underline=True)
        for x in experience:
            title_bits = [x.get('project', ''), x.get('country', ''),
                          str(x.get('year', '') or '')]
            title = ', '.join(b for b in title_bits if b)
            desc  = x.get('description', '')
            if title:
                bullet(title, bold=True)
                for lbl, key in (('Position', 'position'), ('Client', 'client'),
                                 ('Project Description', 'description')):
                    val = x.get(key, '')
                    if val:
                        q = doc.add_paragraph()
                        q.paragraph_format.left_indent = Pt(24)
                        q.paragraph_format.space_after = Pt(2)
                        r = q.add_run(f'{lbl}: ')
                        r.bold = True
                        q.add_run(str(val))
            elif desc:
                # Flat bullet CVs — description-only entries
                bullet(desc)

    # ── Certifications ───────────────────────────────────────────────────────
    certs = engineer.certifications
    if certs:
        section('Professional Courses / Certifications', upper=True)
        for c in certs:
            line = '   '.join(x for x in [c.get('period', ''), c.get('name', '')] if x)
            if line:
                para(line, space_after=1)
            if c.get('location'):
                para(f'Location: {c["location"]}', space_after=6)

    # ── Languages table ──────────────────────────────────────────────────────
    languages = engineer.languages
    if languages:
        section('Languages')
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        for i, h in enumerate(('Languages', 'Reading', 'Speaking', 'Writing')):
            hdr[i].paragraphs[0].add_run(h).bold = True
        for lang in languages:
            row = table.add_row().cells
            row[0].text = str(lang.get('language', ''))
            row[1].text = str(lang.get('reading',  ''))
            row[2].text = str(lang.get('speaking', ''))
            row[3].text = str(lang.get('writing',  ''))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
